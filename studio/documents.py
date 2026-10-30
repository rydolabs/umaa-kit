"""Scoped QEA snapshot -> DOCX *draft* generator.

Outputs model inventories, not official UMAA/ICD specifications. EA diagrams and
Tagged Values are not reconstructed by this module. User templates are local,
trusted DOCX documents. This is not a hardened multiuser upload endpoint.
"""
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
from zipfile import BadZipFile, ZipFile

from .qea import SCHEMA, sha

GENERATOR_VERSION = 'studio-docx-inventory-0.1.0'
MAX_SNAPSHOT_BYTES = 32 * 1024 * 1024
MAX_TEMPLATE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 2000
MAX_CELL_LENGTH = 1500
KINDS = ('packages', 'elements', 'connectors', 'attributes', 'diagrams')
FIELDS = {
    'packages': [('ea_guid', 'EA GUID'), ('Name', 'Package'), ('Parent_ID', 'Parent ID')],
    'elements': [('ea_guid', 'EA GUID'), ('Name', 'Element'), ('Object_Type', 'Type'), ('Package_ID', 'Package ID')],
    'connectors': [('ea_guid', 'EA GUID'), ('Name', 'Connector'), ('Connector_Type', 'Type'),
                   ('Start_Object_ID', 'Source ID'), ('End_Object_ID', 'Target ID')],
    'attributes': [('ea_guid', 'EA GUID'), ('Name', 'Attribute'), ('Type', 'Type'), ('Object_ID', 'Owner ID')],
    'diagrams': [('ea_guid', 'EA GUID'), ('Name', 'Diagram'), ('Diagram_Type', 'Type'), ('Package_ID', 'Package ID')],
}
TABLE_PATTERN = re.compile(r'^\[\[MODEL_TABLE:(packages|elements|connectors|attributes|diagrams)\]\]$')
TEXT_PATTERN = re.compile(r'\{\{MODEL_[A-Z0-9_]+\}\}')


def _load_snapshot(path: Path) -> dict:
    if not path.is_file() or path.stat().st_size > MAX_SNAPSHOT_BYTES:
        raise ValueError('Snapshot missing or exceeds supported file size')
    doc = json.loads(path.read_text(encoding='utf-8'))
    if not isinstance(doc, dict) or doc.get('schemaVersion') != '0.1' or doc.get('sourceKind') != 'sparx-ea-qea':
        raise ValueError('Expected EA snapshot schemaVersion 0.1')
    if doc.get('generationApproved') is not False or doc.get('isReadOnlySnapshot') is not True:
        raise ValueError('Snapshot is not a read-only unapproved EA import')
    if not isinstance(doc.get('sourceSha256'), str) or not re.fullmatch(r'[a-f0-9]{64}', doc['sourceSha256']):
        raise ValueError('Invalid source QEA checksum in snapshot')
    if doc.get('modelValidation') not in ('WARNINGS', 'STRUCTURAL_CHECKS_PASSED'):
        raise ValueError('Unknown structural validation status')
    if not isinstance(doc.get('records'), dict) or not isinstance(doc.get('counts'), dict):
        raise ValueError('Invalid snapshot structure')
    total = 0
    for kind in KINDS:
        rows = doc['records'].get(kind)
        if not isinstance(rows, list) or doc['counts'].get(kind) != len(rows):
            raise ValueError('Snapshot count differs from records: ' + kind)
        total += len(rows)
        seen = set()
        for row in rows:
            if not isinstance(row, dict):
                raise ValueError('Invalid model record')
            guid = row.get('ea_guid')
            # Some EA attribute/diagram exports may omit ea_guid; a draft still
            # can be generated, but no GUID based document approval is claimed.
            if kind in ('packages', 'elements', 'connectors'):
                if not isinstance(guid, str) or not guid or guid in seen:
                    raise ValueError('Missing/duplicate required EA GUID: ' + kind)
                seen.add(guid)
    if total > MAX_ROWS:
        raise ValueError('Snapshot exceeds DOCX inventory row cap; split the model by scope')
    return doc


def _check_template(path: Path):
    if not path.is_file() or path.suffix.lower() != '.docx' or path.stat().st_size > MAX_TEMPLATE_BYTES:
        raise ValueError('Template must be an existing DOCX <= 10 MB')
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
            if len(names) > 2000 or sum(f.file_size for f in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Template archive exceeds decompression limit')
            if any('vbaproject' in n.lower() or n.lower().endswith('.bin') for n in names):
                raise ValueError('Macro/binary parts are not supported in templates')
            if not {'[Content_Types].xml', 'word/document.xml'}.issubset(set(names)):
                raise ValueError('Not a DOCX document')
            # Offline use only; avoid templates with external links when converted
            # by an automated server later.
            for n in names:
                if n.endswith('.rels') and re.search(rb'TargetMode\s*=\s*[\"\']External[\"\']', archive.read(n), re.I):
                    raise ValueError('External template relationships are not supported')
    except BadZipFile as exc:
        raise ValueError('Invalid DOCX template') from exc


def _replace_tokens(paragraph, replacements: dict) -> bool:
    """Replace tokens across runs without discarding adjacent run formatting."""
    runs = paragraph.runs
    joined = ''.join(run.text for run in runs)
    matches = list(TEXT_PATTERN.finditer(joined))
    boundaries = []
    index = 0
    for run in runs:
        boundaries.append((index, index + len(run.text)))
        index += len(run.text)
    for match in reversed(matches):
        token = match.group()
        if token not in replacements:
            raise ValueError('Unknown DOCX placeholder ' + token)
        start, end = match.span()
        left = next(i for i, (a, b) in enumerate(boundaries) if a <= start < b)
        right = next(i for i, (a, b) in enumerate(boundaries) if a < end <= b)
        a = start - boundaries[left][0]
        b = end - boundaries[right][0]
        if left == right:
            runs[left].text = runs[left].text[:a] + replacements[token] + runs[left].text[b:]
        else:
            runs[left].text = runs[left].text[:a] + replacements[token]
            for i in range(left + 1, right):
                runs[i].text = ''
            runs[right].text = runs[right].text[b:]
    if '{{MODEL_' in paragraph.text:
        raise ValueError('Unresolved model placeholder')
    return bool(matches)


def _table(doc, kind: str, rows: list):
    columns = FIELDS[kind]
    t = doc.add_table(rows=1, cols=len(columns))
    try:
        t.style = 'Table Grid'
    except KeyError:
        pass
    for cell, (_, label) in zip(t.rows[0].cells, columns):
        cell.text = label
    for item in rows:
        cells = t.add_row().cells
        for cell, (field, _) in zip(cells, columns):
            value = item.get(field)
            text = '' if value is None else str(value)
            if len(text) > MAX_CELL_LENGTH:
                raise ValueError(f'Model field exceeds report cell limit: {kind}.{field}')
            cell.text = text
    return t


def _make_document(snapshot: dict, template: Path | None, title: str):
    from docx import Document
    from docx.shared import Cm, Pt

    d = Document(str(template)) if template is not None else Document()
    text = {
        '{{MODEL_TITLE}}': title,
        '{{MODEL_SOURCE_SHA256}}': snapshot['sourceSha256'],
        '{{MODEL_VALIDATION}}': snapshot['modelValidation'],
        '{{MODEL_APPROVAL_STATUS}}': 'NOT APPROVED — STRUCTURAL SNAPSHOT ONLY',
    }
    text.update({'{{MODEL_' + kind.upper() + '_COUNT}}': str(snapshot['counts'][kind]) for kind in KINDS})
    if template is None:
        section = d.sections[0]
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        d.add_heading(title, 0)
        d.add_paragraph('DEVELOPMENT DRAFT — NOT UMAA CONFORMANCE OR OFFICIAL ICD APPROVAL')
        d.add_heading('Source and limitations', level=1)
        d.add_paragraph('Original EA QEA SHA-256: ' + snapshot['sourceSha256'])
        d.add_paragraph('Structural import status: ' + snapshot['modelValidation'])
        d.add_paragraph('EA diagrams are listed by metadata only; graphics are not rendered. '
                        'UMAA service semantics, QoS, IDL mapping and normative claims are NOT validated.')
        if snapshot.get('issues'):
            d.add_heading('Structural warnings', level=1)
            for issue in snapshot['issues'][:50]:
                d.add_paragraph(str(issue), style='List Bullet')
        for kind in KINDS:
            d.add_heading(kind.capitalize() + f' ({len(snapshot["records"][kind])})', level=1)
            _table(d, kind, snapshot['records'][kind])
    else:
        found = False
        for p in list(d.paragraphs):
            if TABLE_PATTERN.fullmatch(p.text.strip()):
                kind = TABLE_PATTERN.fullmatch(p.text.strip()).group(1)
                table = _table(d, kind, snapshot['records'][kind])
                p._p.addnext(table._tbl)
                p._element.getparent().remove(p._element)
                found = True
            else:
                found = _replace_tokens(p, text) or found
        for section in d.sections:
            for p in list(section.header.paragraphs) + list(section.footer.paragraphs):
                found = _replace_tokens(p, text) or found
        if not found:
            raise ValueError('Template must contain supported MODEL placeholder or MODEL_TABLE marker')
        # Do not imply that an arbitrary external template is normative.
        d.add_paragraph('Generated from an unapproved EA model snapshot; not UMAA conformance evidence.')
    core = d.core_properties
    core.title = title
    core.subject = 'Unapproved EA QEA development model inventory'
    core.author = 'UMAA Studio (development tool)'
    return d


def generate(snapshot_path: Path, output_path: Path, manifest_path: Path,
             *, template: Path | None = None, title: str = 'EA Model Inventory') -> dict:
    """Create DOCX + matching manifest with no overwrite and minimal partial output."""
    snapshot_path, output_path, manifest_path = map(Path, (snapshot_path, output_path, manifest_path))
    template = Path(template) if template else None
    if not isinstance(title, str) or not title.strip() or len(title) > 160:
        raise ValueError('Invalid report title')
    if output_path.suffix.lower() != '.docx' or manifest_path.suffix.lower() != '.json':
        raise ValueError('Expected a .docx output and .json manifest')
    outputs = [output_path.resolve(), manifest_path.resolve()]
    sources = [snapshot_path.resolve()] + ([template.resolve()] if template else [])
    if len(set(outputs + sources)) != len(outputs + sources):
        raise ValueError('Input and output paths must be distinct')
    if any(p.exists() for p in outputs):
        raise FileExistsError('Refusing to overwrite generated report or manifest')
    snapshot = _load_snapshot(snapshot_path)
    if template:
        _check_template(template)
    doc = _make_document(snapshot, template, title.strip())
    output_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.studio-docx-', dir=output_path.parent) as td:
        staged = Path(td) / 'report.docx'
        doc.save(str(staged))
        report_hash = sha(staged)
        manifest = {
            'schemaVersion': '1.0', 'kind': 'studio-ea-model-inventory',
            'generatorVersion': GENERATOR_VERSION,
            'generatedAtUtc': datetime.now(timezone.utc).isoformat(),
            'modelSourceKind': 'sparx-ea-qea',
            'sourceQeaSha256AsDeclaredBySnapshot': snapshot['sourceSha256'],
            'snapshotSha256': sha(snapshot_path),
            'templateSha256': sha(template) if template else None,
            'docxSha256': report_hash,
            'rowCounts': snapshot['counts'],
            'modelStructuralValidation': snapshot['modelValidation'],
            'approved': False, 'conformanceClaim': False,
            'limitations': ['Not a formal UMAA specification', 'EA diagrams not rendered',
                            'Generated from an unapproved QEA model snapshot'],
        }
        # Exclusive output creation; manifest is never exposed before DOCX bytes
        # are committed. A failed manifest write removes only our created DOCX.
        with output_path.open('xb') as f, staged.open('rb') as src:
            import shutil
            shutil.copyfileobj(src, f)
        try:
            with manifest_path.open('x', encoding='utf-8') as f:
                json.dump(manifest, f, indent=2, ensure_ascii=False)
                f.write('\n')
        except BaseException:
            output_path.unlink(missing_ok=True)
            raise
    return manifest


def verify_document(snapshot_path: Path, docx_path: Path, manifest_path: Path,
                    template: Path | None = None) -> list[str]:
    """Consistency check; SHA-256 integrity is not origin authenticity."""
    errors = []
    try:
        doc = json.loads(Path(manifest_path).read_text(encoding='utf-8'))
        if (doc.get('schemaVersion') != '1.0' or doc.get('kind') != 'studio-ea-model-inventory'
                or doc.get('generatorVersion') != GENERATOR_VERSION):
            errors.append('Unsupported manifest schema or generator version')
        if doc.get('approved') is not False or doc.get('conformanceClaim') is not False:
            errors.append('Unsafe document approval/conformance claim')
        snapshot = _load_snapshot(Path(snapshot_path))
        for key, expected in (
            ('snapshotSha256', sha(Path(snapshot_path))),
            ('sourceQeaSha256AsDeclaredBySnapshot', snapshot['sourceSha256']),
            ('docxSha256', sha(Path(docx_path))),
            ('templateSha256', sha(Path(template)) if template else None),
            ('rowCounts', snapshot['counts']),
            ('modelStructuralValidation', snapshot['modelValidation']),
        ):
            if doc.get(key) != expected:
                errors.append('Manifest mismatch: ' + key)
        with ZipFile(docx_path) as archive:
            if archive.testzip() is not None or 'word/document.xml' not in archive.namelist():
                errors.append('Generated DOCX archive is invalid')
    except (OSError, ValueError, KeyError, TypeError, BadZipFile) as exc:
        errors.append('Document evidence cannot be verified: ' + str(exc))
    return errors
