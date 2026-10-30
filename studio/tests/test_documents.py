"""Synthetic QEA-like SQLite source fixtures; NOT real Sparx acceptance tests."""
import copy
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from studio.documents import generate, verify_document
from studio.qea import extract, save_snapshot
from test_qea import qea_fixture


class DocumentGenerationTests(unittest.TestCase):
    def _snapshot(self, root):
        src = qea_fixture(root / 'sample.qea')
        snapshot = root / 'model.json'
        save_snapshot(extract(src), snapshot)
        return src, snapshot

    def test_generated_report_and_manifest(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            src, snap = self._snapshot(root)
            docx, meta = root / 'inventory.docx', root / 'inventory.json'
            before = src.read_bytes()
            res = generate(snap, docx, meta, title='Engine Inventory')
            self.assertEqual(before, src.read_bytes())
            self.assertFalse(res['conformanceClaim'])
            self.assertFalse(res['approved'])
            self.assertEqual(verify_document(snap, docx, meta), [])
            text = '\n'.join(p.text for p in Document(docx).paragraphs)
            self.assertIn('Engine Inventory', text)
            self.assertIn('NOT UMAA CONFORMANCE', text)
            cells = [c.text for t in Document(docx).tables for r in t.rows for c in r.cells]
            self.assertIn('Engine', cells)
            self.assertIn('obj-7', cells)
            with self.assertRaises(FileExistsError):
                generate(snap, docx, meta)

    def test_template_with_split_runs_and_tables(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            _, snap = self._snapshot(root)
            tpl = root / 'template.docx'
            t = Document()
            p = t.add_paragraph()
            p.add_run('Model: ').bold = True
            p.add_run('{{MODEL_')
            p.add_run('TITLE}}')
            p.add_run(' — {{MODEL_ELEMENTS_COUNT}} elements')
            t.add_paragraph('[[MODEL_TABLE:elements]]')
            t.add_paragraph('SHA {{MODEL_SOURCE_SHA256}}')
            t.save(tpl)
            docx, meta = root / 'r.docx', root / 'r.json'
            generate(snap, docx, meta, template=tpl, title='Verification Preview')
            doc = Document(docx)
            full_text = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('Verification Preview', full_text)
            self.assertIn('2 elements', full_text)
            self.assertIn('SHA ', full_text)
            self.assertNotIn('MODEL_', full_text)
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(len(doc.tables[0].rows), 3)
            self.assertEqual(verify_document(snap, docx, meta, tpl), [])

    def test_unknown_template_token_rejected_no_artifact(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            _, snap = self._snapshot(root)
            tpl = root / 't.docx'
            d = Document()
            d.add_paragraph('{{MODEL_MYSTERY}}')
            d.save(tpl)
            with self.assertRaisesRegex(ValueError, 'Unknown DOCX placeholder'):
                generate(snap, root / 'a.docx', root / 'm.json', template=tpl)
            self.assertFalse((root / 'a.docx').exists())

    def test_metadata_forgery_and_record_count_mismatch_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            _, snap = self._snapshot(root)
            doc = json.loads(snap.read_text())
            doc['generationApproved'] = True
            bad = root / 'forged.json'
            bad.write_text(json.dumps(doc))
            with self.assertRaisesRegex(ValueError, 'unapproved'):
                generate(bad, root / 'a.docx', root / 'a.json')
            doc['generationApproved'] = False
            doc['counts']['elements'] = 99
            bad.write_text(json.dumps(doc))
            with self.assertRaisesRegex(ValueError, 'count differs'):
                generate(bad, root / 'a.docx', root / 'a.json')

    def test_tamper_docx_and_manifest_detected(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            _, snap = self._snapshot(root)
            out, meta = root / 'a.docx', root / 'a.json'
            generate(snap, out, meta)
            with out.open('ab') as stream:
                stream.write(b'attack')
            self.assertTrue(any('docxSha256' in e for e in verify_document(snap, out, meta)))
            out.unlink()
            meta.unlink()
            generate(snap, out, meta)
            doc = json.loads(meta.read_text())
            doc['conformanceClaim'] = True
            meta.write_text(json.dumps(doc))
            self.assertTrue(any('Unsafe document' in e for e in verify_document(snap, out, meta)))
            doc['conformanceClaim'] = False
            doc['modelStructuralValidation'] = 'FALSIFIED_PASS'
            meta.write_text(json.dumps(doc))
            self.assertTrue(any('modelStructuralValidation' in e for e in verify_document(snap, out, meta)))

    def test_cell_size_limit_blocks_without_partial_file(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            _, snap = self._snapshot(root)
            data = json.loads(snap.read_text())
            data['records']['elements'][0]['Name'] = 'Z' * 1501
            snap.write_text(json.dumps(data))
            with self.assertRaisesRegex(ValueError, 'cell limit'):
                generate(snap, root / 'a.docx', root / 'a.json')
            self.assertFalse((root / 'a.docx').exists())


if __name__ == '__main__':
    unittest.main()
